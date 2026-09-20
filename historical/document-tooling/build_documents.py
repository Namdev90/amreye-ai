from pathlib import Path
from copy import deepcopy
import re,json,hashlib,collections
from io import BytesIO
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).parent
OUT=ROOT.parent/'outputs'
SUP=OUT/'Supporting documents'
OUT.mkdir(exist_ok=True);SUP.mkdir(exist_ok=True)
SOURCE=ROOT/'source-master.docx'
source=Document(SOURCE)
blocks=json.loads((ROOT/'source-blocks.json').read_text(encoding='utf-8'))
byidx={x['index']:x for x in blocks}
modules=[
 {'key':'P1','file':'AMREYE AI Product and Laboratory Workflows.docx','title':'AMREYE AI Product and Laboratory Workflows','ranges':[(19,464),(2800,3486)],'intro':'Detailed product functions, laboratory workflows, intended applications and expansion boundaries. The immediate product is Reader V1 for endpoint disk diffusion. Broader sectors and capabilities remain proposed or subject to their own evidence.','names':['Project scope and requirements','Product portfolio and laboratory workflows','Intended applications and expansion boundaries']},
 {'key':'P2','file':'AMREYE AI Hardware and Sensing Engineering.docx','title':'AMREYE AI Hardware and Sensing Engineering','ranges':[(465,1241)],'intro':'Engineering reference for imaging, calibration, sensors and instrument configuration. The listed sensing options are design and research candidates; their inclusion does not establish an implemented instrument or a validated biological measurement.','names':['Sensing engineering and technical design']},
 {'key':'P3','file':'AMREYE AI Software and Data Systems.docx','title':'AMREYE AI Software and Data Systems','ranges':[(1242,2799),(6571,7029)],'intro':'Detailed AI functions, software architecture, interoperability, plate records, digital twins and data definitions. Statements describing platform functions are design intent unless accompanied by project-specific implementation evidence. Digital records, simulations and measured results remain distinct.','names':['AI software data and interoperability','Digital plate records twins and simulation','Product data contracts and terminology']},
 {'key':'P4','file':'AMREYE AI Validation Legal and Safety.docx','title':'AMREYE AI Validation Legal and Safety','ranges':[(3487,4313)],'intro':'Validation, claim interpretation, legal and regulatory considerations, safety and security requirements. This is the retained project planning reference; it does not establish approval, registration, completed legal review or successful validation.','names':['Validation claim control safety and security']},
 {'key':'P5','file':'AMREYE AI Scientific Evidence and References.docx','title':'AMREYE AI Scientific Evidence and References','ranges':[(4314,5415),(7030,7424)],'intro':'Research assessment, technical audit and the complete retained reference catalogue. External studies provide evidence about their own methods and conditions; they do not establish AMREYE performance. Reference counts are inventories, not counts of independent confirmation.','names':['Scientific evidence and research assessment','Citation statistics and research coverage','References and source register']},
 {'key':'P6','file':'AMREYE AI Development Roadmap and Advanced Architecture.docx','title':'AMREYE AI Development Roadmap and Advanced Architecture','ranges':[(5416,6429)],'intro':'Reader V1 development priorities, implementation planning, and proposed Bio-Vault and APEX XR configurations. Advanced specifications, interfaces and subsystem descriptions remain proposed architecture. Dates, durations and capability descriptions do not establish delivery commitments or completed functions.','names':['Development roadmap and advanced configurations']},
 {'key':'P7','file':'AMREYE AI Bill of Materials and Financial Planning.docx','title':'AMREYE AI Bill of Materials and Financial Planning','ranges':[(6430,6570)],'intro':'Component and procurement detail, working budgets, illustrative pricing and scenario models. Figures are internal planning inputs or dated observations. Reader V1, integrated Bio-Vault and APEX XR scopes must remain separate. The source records an unresolved INR 500 difference between the summary BOM and detailed component sum.','names':['Financial planning and bill of materials']},
]
owners={}
for m in modules:
 for lo,hi in m['ranges']:
  for i in range(lo,hi+1):
   assert i not in owners
   owners[i]=m['key']
assert set(owners)==set(range(19,7425))

def clean_text(t):
 return t.replace('\u2014','; ').replace('\u2013','-')

def add_link(p,text,target,anchor=None):
 h=OxmlElement('w:hyperlink')
 if anchor: h.set(qn('w:anchor'),anchor)
 else: h.set(qn('r:id'),p.part.relate_to(target,RT.HYPERLINK,is_external=True))
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
 col=OxmlElement('w:color');col.set(qn('w:val'),'154F63');pr.append(col)
 under=OxmlElement('w:u');under.set(qn('w:val'),'single');pr.append(under)
 r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)

def bookmark(p,name,bid):
 a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(bid));a.set(qn('w:name'),name)
 z=OxmlElement('w:bookmarkEnd');z.set(qn('w:id'),str(bid));p._p.insert(1 if p._p.pPr is not None else 0,a);p._p.append(z)

def configure(d,header,main=False):
 for sec in d.sections:
  sec.page_width=Inches(8.5);sec.page_height=Inches(14)
  sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(.5)
  sec.header_distance=sec.footer_distance=Inches(.22)
  for p in sec.header.paragraphs:p.clear()
  hp=sec.header.paragraphs[0];hp.text=header;hp.style=d.styles['Header']
  hp.runs[0].font.size=Pt(8);hp.runs[0].font.color.rgb=RGBColor(0,0,0)
  for p in sec.footer.paragraphs:p.clear()
  p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
  p.add_run('AMREYE.AI  |  ')
  f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
  for r in p.runs:r.font.size=Pt(8)
 for s in d.styles:
  if s.type==1:
   s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
   if s.name.startswith(('Heading','Title','Subtitle','Header')):
    for c in s.element.xpath('.//w:color'):
     for a in list(c.attrib):
      if a!=qn('w:val'):del c.attrib[a]
     c.set(qn('w:val'),'000000')
    for border in s.element.xpath('.//w:pBdr'):border.getparent().remove(border)
 n=d.styles['Normal'];n.font.size=Pt(11 if main else 10)
 n.paragraph_format.space_after=Pt(6 if main else 2.5)
 n.paragraph_format.line_spacing=1.12 if main else 1.04
 for name,size,before,after in [('Title',25,0,10),('Heading 1',17,15,7),('Heading 2',12.5,10,5),('Heading 3',11,7,3)]:
  s=d.styles[name];s.font.size=Pt(size);s.font.bold=True
  s.paragraph_format.space_before=Pt(before);s.paragraph_format.space_after=Pt(after)
  s.paragraph_format.keep_with_next=True;s.paragraph_format.page_break_before=False
 d.styles['Caption'].font.size=Pt(9)
 d.styles['Caption'].font.color.rgb=RGBColor.from_string('333333')
 d.styles['Caption'].paragraph_format.keep_with_next=True

def style_table(t,main=False):
 t.alignment=WD_TABLE_ALIGNMENT.CENTER
 t.autofit=False
 labels=[c.text.strip() for c in t.rows[0].cells]
 widths=None
 if len(labels)==4 and labels[0]=='ID':widths=[.45,4.5,.5,2.05]
 elif len(labels)==3 and labels[0]=='BOM category':widths=[3.1,.75,3.65]
 if widths:
  for c,w in zip(t.columns,widths):c.width=Inches(w)
  for row in t.rows:
   for c,w in zip(row.cells,widths):c.width=Inches(w)
 props=t._tbl.tblPr
 for old in props.findall(qn('w:tblBorders')):props.remove(old)
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
 props.insert_element_before(borders,'w:shd','w:tblLayout','w:tblCellMar','w:tblLook','w:tblCaption','w:tblDescription','w:tblPrChange')
 for ri,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr()
  for h in list(trpr.findall(qn('w:trHeight'))):trpr.remove(h)
  if trpr.find(qn('w:cantSplit')) is None:trpr.append(OxmlElement('w:cantSplit'))
  if ri==0:
   for old in list(trpr.findall(qn('w:tblHeader'))):trpr.remove(old)
   h=OxmlElement('w:tblHeader');trpr.append(h)
  for cell in row.cells:
   cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=cell._tc.get_or_add_tcPr()
   mar=OxmlElement('w:tcMar')
   for e in ['top','left','bottom','right']:
    v=OxmlElement('w:'+e);v.set(qn('w:w'),'85' if e in ['top','bottom'] else '100');v.set(qn('w:type'),'dxa');mar.append(v)
   for old in list(pr.findall(qn('w:tcMar'))):pr.remove(old)
   pr.insert_element_before(mar,'w:textDirection','w:tcFitText','w:vAlign','w:hideMark','w:headers','w:cellIns','w:cellDel','w:cellMerge','w:tcPrChange')
   for sh in list(pr.findall(qn('w:shd'))):pr.remove(sh)
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'173F51' if ri==0 else ('F0F4F6' if ri%2==0 else 'FFFFFF'));pr.insert_element_before(sh,'w:noWrap','w:tcMar','w:textDirection','w:tcFitText','w:vAlign','w:hideMark','w:headers','w:cellIns','w:cellDel','w:cellMerge','w:tcPrChange')
   for p in cell.paragraphs:
    p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.06
    p.paragraph_format.keep_with_next=(ri==0)
    for r in p.runs:
     r.font.name='Calibri';r.font.size=Pt(10 if main else 9)
     r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
     if ri==0:r.bold=True

def copy_companion(m):
 d=Document()
 d.part._styles_part._element=deepcopy(source.styles.element)
 d.part.numbering_part._element=deepcopy(source.part.numbering_part.element)
 body=d.element.body
 sect=deepcopy(body.sectPr)
 for child in list(body):body.remove(child)
 body.append(sect)
 configure(d,m['title'].replace('AMREYE AI ','')+'  |  Supporting reference')
 d.add_paragraph(m['title'],'Title')
 d.add_paragraph(m['intro'])
 d.add_paragraph('Supporting reference revised 19 September 2026 from the source master last updated 17 September 2026. Detailed source content is retained for specialist use. Original figure and table numbers are retained for traceability. The concise AMREYE.AI Super Document remains the main project overview.','Normal')
 if m['key']=='P7':d.add_paragraph('The INR 10 lakh programme allocation is for first core Reader V1 development, research and validation. The reference component BOM and the later integrated-platform and APEX XR scenarios have separate scopes.')
 p=d.add_paragraph();add_link(p,'Open the main project overview','../AMREYE.AI Super Document.docx')
 d.add_paragraph('Contents','Heading 2')
 for j,name in enumerate(m['names']):
  p=d.add_paragraph();add_link(p,name,None,anchor=f'{m["key"]}_chapter_{j+1}')
 if m['key']!='P5':
  p=d.add_paragraph();add_link(p,'Full source catalogue','AMREYE AI Scientific Evidence and References.docx')
 d.add_paragraph()
 included=[]
 chapter_no=0
 for lo,hi in m['ranges']:
  for i in range(lo,hi+1):
   el=deepcopy(source.element.body[i])
   for element in el.iter():
    for attr,rid in list(element.attrib.items()):
     if attr.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'):
      rel=source.part.rels[rid]
      if rel.reltype==RT.IMAGE:
       newrid,_=d.part.get_or_add_image(BytesIO(rel.target_part.blob))
      elif rel.is_external:newrid=d.part.relate_to(rel.target_ref,rel.reltype,is_external=True)
      else:raise ValueError('Unexpected imported relationship '+rel.reltype)
      element.set(attr,newrid)
   body.insert(len(body)-1,el);included.append(i)
   if i==36:
    from docx.text.paragraph import Paragraph
    p=Paragraph(el,d)
    p.text='Qualified plate and method → controlled imaging → readability and quality checks → calibrated zone measurement → versioned standards interpretation → human review → report or laboratory and surveillance export.'
    d.add_paragraph('Image or model performance, dimensional measurement performance and categorical AST agreement are different outcomes. Synthetic demonstrations and external benchmarks do not establish AMREYE performance without project-specific validation.')
   if i==2217:
    from docx.text.paragraph import Paragraph
    p=Paragraph(el,d)
    p.text='Proposed architecture: retain a canonical plate and result model containing raw measurements, method, standards version, quality checks and reviewer provenance. A versioned adapter layer manages site mappings, terminology, validation and audit.'
    t=d.add_table(rows=1,cols=2)
    t.rows[0].cells[0].text='Proposed destination';t.rows[0].cells[1].text='Role in the architecture'
    for target,purpose in [('FHIR R4 and ABDM','DiagnosticReport and Observation representations'),('HL7 v2 and site LIS','Legacy and enterprise interfaces'),('AST analyser links','Connections required by particular instruments or sites'),('WHONET and BacLink','Local analytics and GLASS-ready export'),('OpenELIS Global','Reference integration target rather than the product architecture'),('National and regional systems','Programme-specific adapters, including NARS-Net, EARS-Net, NHSN and ReLAVRA+')]:
     cells=t.add_row().cells;cells[0].text=target;cells[1].text=purpose
    d.add_paragraph('The design keeps the canonical record stable while adapting outward by site, jurisdiction and surveillance programme. Measurements remain reproducible as adapters and external standards change; the instrument does not require direct firmware support for every destination.')
   if byidx.get(i,{}).get('style')=='Heading 1':
    from docx.text.paragraph import Paragraph
    oldmarks=[deepcopy(x) for x in el.xpath('./w:bookmarkStart|./w:bookmarkEnd')]
    p=Paragraph(el,d);p.text=m['names'][chapter_no]
    for mark in oldmarks:el.append(mark)
    bookmark(p,f'{m["key"]}_chapter_{chapter_no+1}',200000+chapter_no)
    chapter_no+=1
   if i==389:
    from docx.text.paragraph import Paragraph
    p=Paragraph(el,d);p.text='Detailed financial assumptions and scenario models are retained in the Bill of Materials and Financial Planning companion.'
   if i==6546:
    from docx.text.paragraph import Paragraph
    p=Paragraph(el,d);p.style=d.styles['Normal'];p.text='Detailed flagship architecture is covered in the Development Roadmap and Advanced Architecture companion. The following table retains the corresponding BOM categories and pricing assumptions.'
 for p in d.paragraphs:
  if p.style.name in ['Normal','List Paragraph']:
   p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(2.5);p.paragraph_format.line_spacing=1.04
   if p.text.startswith('Price code:'):p.paragraph_format.keep_with_next=True
  if p.style.name.startswith(('Heading','Title')):
   p.paragraph_format.keep_with_next=True;p.paragraph_format.keep_together=True
   for r in p.runs:
    r.font.color.rgb=RGBColor(0,0,0)
    for c in r._r.xpath('.//w:color'):
     for a in list(c.attrib):
      if a!=qn('w:val'):del c.attrib[a]
  # Existing source font sizes are harmonized without changing emphasis or wording.
  if p.style.name in ['Normal','List Paragraph']:
   for r in p.runs:r.font.size=Pt(10);r.font.name='Calibri'
   visible=[r for r in p.runs if r.text.strip()]
   if visible and len(p.text)<180 and all(r.bold for r in visible):p.paragraph_format.keep_with_next=True
   prefixes=('Price code:','Mechanical function:','Bridge item:','Workstream:','Question:','Risk:','Priority:','Layer:','Level:','Capability:','Proposed question:','Evidence path:','Comparison dimension:','Candidate alternative:','Cost basis:','Event:','Safe current claim:','Hardware maturity:')
   if len(p.text)<150 and p.text.startswith(prefixes):p.paragraph_format.keep_with_next=True
 for t in d.tables:style_table(t)
 # A paragraph must follow a final table before the section properties.
 if len(body)>1 and body[-2].tag==qn('w:tbl'):d.add_paragraph()
 d.core_properties.title=m['title'];d.core_properties.subject=m['intro'][:250]
 d.core_properties.author='AMREYE.AI';d.core_properties.comments='Detailed source content redistributed from the canonical master; claims retain their original evidence status.'
 path=SUP/m['file'];d.save(path)
 return included

def inline(p,text):
 for j,t in enumerate(re.split(r'(\*\*[^*]+\*\*)',text)):
  if t.startswith('**') and t.endswith('**'):p.add_run(t[2:-2]).bold=True
  else:p.add_run(t)

def build_main():
 d=Document();configure(d,'AMREYE.AI  |  Project overview',main=True)
 lines=(ROOT/'main.md').read_text(encoding='utf-8').splitlines()
 i=0;heading_no=0
 while i<len(lines):
  line=lines[i].strip()
  if not line:i+=1;continue
  if line.startswith('|'):
   rows=[]
   while i<len(lines) and lines[i].strip().startswith('|'):
    cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
    if not all(re.fullmatch(r'[-: ]+',x) for x in cells):rows.append(cells)
    i+=1
   t=d.add_table(rows=0,cols=len(rows[0]));widths=[2.0,5.5] if len(rows[0])==2 else [1.45,3.7,2.35]
   if len(rows[0])==2 and 'allocation' in rows[0][1].lower():widths=[5.8,1.7]
   for col,w in zip(t.columns,widths):col.width=Inches(w)
   for row in rows:
    cells=t.add_row().cells
    for ci,tx in enumerate(row):
     cells[ci].width=Inches(widths[ci]);inline(cells[ci].paragraphs[0],tx)
   style_table(t,True);d.add_paragraph();continue
  if line.startswith('!['):
   path=line.split('](',1)[1][:-1];d.add_picture(str(ROOT/path),width=Inches(7.35))
   d.add_paragraph('Reader V1 proposed workflow  |  Conceptual diagram retained from the source master','Caption')
  elif line.startswith('# '):d.add_paragraph(line[2:],'Title')
  elif line.startswith('## '):
   p=d.add_paragraph(line[3:],'Heading 1');heading_no+=1;bookmark(p,f'main_section_{heading_no}',300000+heading_no)
  elif line.startswith('### '):d.add_paragraph(line[4:],'Heading 2')
  else:
   p=d.add_paragraph();inline(p,line)
  i+=1
 # Lightweight linked contents under the opening date, no separate contents page.
 headings=[p.text for p in d.paragraphs if p.style.name=='Heading 1']
 anchor=d.paragraphs[2]
 p=d.add_paragraph();p._p.getparent().remove(p._p);anchor._p.addnext(p._p)
 p.add_run('In this document  ').bold=True
 for j,h in enumerate(headings):
  if j:p.add_run('  ·  ')
  add_link(p,re.sub(r'^\d+\s+','',h),None,anchor=f'main_section_{j+1}')
 for r in p.runs:r.font.size=Pt(9)
 # Add clickable companion references to the library rows and internal references.
 table=d.tables[-1]
 for m,row in zip(modules,table.rows[1:]):
  p=row.cells[0].paragraphs[0];label=p.text;p.clear();add_link(p,label,'Supporting documents/'+m['file'])
 for p in d.paragraphs:
  if p.style.name=='Normal':p.paragraph_format.keep_together=True
  match=re.match(r'P([1-7])\. ',p.text)
  if match:
   m=modules[int(match.group(1))-1];txt=p.text;p.clear();add_link(p,txt,'Supporting documents/'+m['file'])
  urls=re.findall(r'https?://\S+',p.text)
  if urls:
   txt=p.text;p.clear()
   for u in urls:
    before,txt=txt.split(u,1);p.add_run(before);add_link(p,u,u)
   p.add_run(txt)
 d.core_properties.title='AMREYE AI Project Overview';d.core_properties.author='AMREYE.AI'
 d.core_properties.subject='Concise project overview with supporting specialist references'
 d.save(OUT/'AMREYE.AI Super Document.docx')

def fix_links():
 # Move source navigation targets across companions without leaving broken anchors.
 bookmark_owner={}
 for m in modules:
  d=Document(SUP/m['file'])
  for b in d.element.body.xpath('.//w:bookmarkStart'):
   bookmark_owner[b.get(qn('w:name'))]=m['file']
 fixed=0;unresolved=[]
 for m in modules:
  path=SUP/m['file'];d=Document(path)
  local={b.get(qn('w:name')) for b in d.element.body.xpath('.//w:bookmarkStart')}
  for h in d.element.body.xpath('.//w:hyperlink'):
   anchor=h.get(qn('w:anchor'))
   if anchor and anchor not in local:
    if anchor in bookmark_owner:
     h.set(qn('r:id'),d.part.relate_to(bookmark_owner[anchor]+'#'+anchor,RT.HYPERLINK,is_external=True));h.attrib.pop(qn('w:anchor'));fixed+=1
    else:
     # Source-level navigation with no target in substantive content becomes plain text.
     parent=h.getparent();idx=parent.index(h)
     for ch in list(h):parent.insert(idx,ch);idx+=1
     parent.remove(h);unresolved.append(anchor)
  d.save(path)
 return {'cross_document_links':fixed,'obsolete_source_navigation_removed':sorted(set(unresolved))}

if __name__=='__main__':
 import sys
 if '--main-only' in sys.argv:build_main();print('Main document built');raise SystemExit
 mapping={}
 for m in modules:mapping[m['key']]=copy_companion(m)
 linkreport=fix_links()
 manifest={'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'source_words':sum(len(x['text'].split())for x in blocks),'coverage':'Every substantive source body block 19 through 7424 assigned exactly once','modules':modules,'mapping':mapping,'link_report':linkreport,'editorial_changes':[{'block':389,'change':'Replace obsolete financial-exclusion statement with pointer to finance companion'},{'block':6546,'change':'Repair fragment as complete sentence with pointer to architecture companion'},{'block':36,'change':'Re-express workflow illustration and its qualification as editable text; original artwork remains in source backup'},{'block':2217,'change':'Re-express crowded adapter architecture illustration as editable prose and mapping table; original artwork remains in source backup'}],'navigation_removed':'Original cover, global contents and global figures/tables list replaced with document-specific navigation'}
 (ROOT/'distribution-manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
 print(json.dumps({'companions':len(modules),'source_blocks_assigned':len(owners),'links':linkreport}))
