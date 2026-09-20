from docx import Document
from pathlib import Path
import json,re,zipfile
root=Path(__file__).parent
d=Document(root/'source-master.docx')
blocks=[]
for i,el in enumerate(d.element.body):
    typ=el.tag.split('}')[-1]
    text=' | '.join(''.join(p.itertext()) for p in [])
    if typ=='p':
        from docx.text.paragraph import Paragraph
        p=Paragraph(el,d)
        blocks.append({'index':i,'type':'p','style':p.style.name,'text':p.text})
    elif typ=='tbl':
        from docx.table import Table
        t=Table(el,d)
        rows=[[c.text for c in r.cells] for r in t.rows]
        blocks.append({'index':i,'type':'table','rows':rows,'text':'\n'.join(' | '.join(r) for r in rows)})
(root/'source-blocks.json').write_text(json.dumps(blocks,ensure_ascii=False,indent=2),encoding='utf-8')
(root/'source-text.txt').write_text('\n\n'.join(f"[{b['index']}] {b.get('style',b['type'])}\n{b['text']}" for b in blocks),encoding='utf-8')
headings=[{'index':b['index'],'style':b.get('style'),'text':b['text']} for b in blocks if b.get('style','').startswith(('Heading','Title'))]
(root/'source-headings.json').write_text(json.dumps(headings,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'words':sum(len(b['text'].split()) for b in blocks),'paragraphs':len(d.paragraphs),'tables':len(d.tables),'images':len(d.inline_shapes),'headings':headings,'sections':[{'width':round(s.page_width.inches,2),'height':round(s.page_height.inches,2),'margins':[round(s.left_margin.inches,2),round(s.right_margin.inches,2)]} for s in d.sections]},ensure_ascii=False))
